import os
import re
import argparse
import PyPDF2
from docx import Document
from odf.opendocument import load as load_odt
import xlrd
import openpyxl
import pandas as pd

def count_words(text):
    return len(re.findall(r'\w+', text))

def get_text_from_odf_element(element):
    texts = []
    if hasattr(element, 'textContent') and element.textContent:
        texts.append(element.textContent)
    elif hasattr(element, 'childNodes'):
        for node in element.childNodes:
            if node.nodeType == node.TEXT_NODE:
                texts.append(node.data)
            elif node.nodeType == node.ELEMENT_NODE:
                texts.append(get_text_from_odf_element(node))
    return ' '.join(texts)

def count_words_in_txt(path):
    with open(path, 'r', encoding='utf-8') as file:
        text = file.read()
    return count_words(text)

def count_words_in_pdf(path):
    with open(path, 'rb') as file:
        reader = PyPDF2.PdfReader(file)
        text = ''
        for page in reader.pages:
            text += page.extract_text()
    return count_words(text)

def count_words_in_docx(path):
    doc = Document(path)
    text = ''
    for p in doc.paragraphs:
        text += p.text
    return count_words(text)

def count_words_in_odt(path):
    doc = load_odt(path)
    texts = []
    for node in doc.childNodes:
        if node.nodeType == node.TEXT_NODE:
            texts.append(node.data)
        elif node.nodeType == node.ELEMENT_NODE:
            texts.append
    text = ' '.join(texts)
    return count_words(text)

def count_words_in_xlsx(path):
    wb = openpyxl.load_workbook(path)
    sheet = wb.active
    text = ''
    for row in sheet.iter_rows(values_only=True):
        for cell in row:
            text += str(cell) + ' '
    return count_words(text)

def count_words_in_xls(path):
    wb = xlrd.open_workbook(path)
    sheet = wb.sheet_by_index(0)
    text = ''
    for row in range(sheet.nrows):
        for col in range(sheet.ncols):
            text += str(sheet.cell_value(row, col)) + ' '
    return count_words(text)

def count_words_in_csv(path):
    df = pd.read_csv(path)
    text = df.to_string(index=False)
    return count_words(text)

def count_words_in_ods(path):
    doc = load_odt(path)
    text = get_text_from_odf_element(doc)
    return count_words(text)

def count_words_in_file(path):
    _, extension = os.path.splitext(path)
    if extension == '.pdf':
        return count_words_in_pdf(path)
    elif extension == '.docx':
        return count_words_in_docx(path)
    elif extension == '.odt':
        return count_words_in_odt(path)
    elif extension == '.xlsx':
        return count_words_in_xlsx(path)
    elif extension == '.xls':
        return count_words_in_xls(path)
    elif extension == '.ods':
        return count_words_in_ods(path)
    elif extension == '.csv':
        return count_words_in_csv(path)
    else:
        return count_words_in_txt(path)

argparser = argparse.ArgumentParser()
argparser.add_argument('--input', '-i', help='input file path', required=True)
args = argparser.parse_args()

print(count_words_in_file(args.input))
